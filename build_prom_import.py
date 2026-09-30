#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_prom_import.py
=====================

Builds a Prom.ua XLS(X) import file using the PROVEN update pattern:

  1. Start from a full product export you already downloaded from Prom.ua
     (Товари -> Експорт). This file already has every product's own real
     name (RU + UK), description, images, characteristics, and — crucially
     — its "Унікальний_ідентифікатор" (Prom's permanent internal product
     ID), which is what Prom actually uses to match a row in an import
     file to an existing product.
  2. Download the supplier feeds and, for every product whose "Код_товару"
     (SKU/article) matches a supplier's SKU, overwrite ONLY the
     Ціна / Наявність / Кількість cells for that row with the fresh
     supplier data.
  3. Leave every other cell — name, description, translations, images,
     characteristics, group, unique id — exactly as it was in the export.

Why this instead of the YML "update via link" feed:
Prom's automatic YML update matches products by the `id` attribute of
`<offer>`, which must equal a product ID Prom itself assigned — not the
supplier's SKU. Products that were never created through that exact YML
feed don't have this ID assigned, so Prom can't recognize them and creates
duplicates instead of updating. The XLS/CSV import route matches by the
"Унікальний_ідентифікатор" column, which is ALWAYS present on any export,
so it reliably updates existing products without duplicating or touching
their name/translations (as long as you don't touch the Ідентифікатор
columns and leave the *_укр fields exactly as exported — Prom skips
touching a translation entirely if its укр fields are left blank/unchanged
here, since we simply reuse whatever was already exported).

Usage:
    python build_prom_import.py --base export-products.xlsx --output prom_import.xlsx

Run with --inspect first if you're not sure the column names in your
export match what this script expects (see HEADER_VARIANTS below).
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path
from typing import Optional

import openpyxl

import merge_feeds as feeds  # reuses the supplier-feed download/parse logic

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger("build_prom_import")

# --------------------------------------------------------------------------
# Header name variants. Prom exports column headers in whatever language
# the account's interface is set to, so the exact same field can show up
# as "Код_товару" (Ukrainian) or "Код_товара" (Russian) depending on the
# account. We match case-insensitively against every variant below.
# --------------------------------------------------------------------------
HEADER_VARIANTS = {
    "sku": ["Код_товару", "Код_товара"],
    "price": ["Ціна", "Цена"],
    "available": ["Наявність", "Наличие"],
    "quantity": ["Кількість", "Количество"],
    "unique_id": ["Унікальний_ідентифікатор", "Уникальный_идентификатор"],
}

SHEET_NAME = "Export Products Sheet"


def find_column(headers: list[Optional[str]], variants: list[str]) -> Optional[int]:
    """Returns the 1-based column index matching one of the given header
    name variants (case-insensitive), or None if not found."""
    lower_variants = {v.lower() for v in variants}
    for idx, h in enumerate(headers, start=1):
        if h and str(h).strip().lower() in lower_variants:
            return idx
    return None


def load_column_map(ws) -> dict[str, int]:
    headers = [c.value for c in ws[1]]
    col_map = {}
    for field, variants in HEADER_VARIANTS.items():
        idx = find_column(headers, variants)
        if idx is None:
            log.warning(
                "Could not find a column for '%s' (looked for: %s). "
                "This field will not be updated.",
                field, variants,
            )
        else:
            col_map[field] = idx
    return col_map


def update_workbook(base_path: Path, output_path: Path, mode: str = "both") -> None:
    log.info("Loading base export: %s", base_path)
    wb = openpyxl.load_workbook(base_path)
    if SHEET_NAME not in wb.sheetnames:
        log.error(
            "Sheet '%s' not found in the base file. Found sheets: %s",
            SHEET_NAME, wb.sheetnames,
        )
        sys.exit(1)
    ws = wb[SHEET_NAME]

    col_map = load_column_map(ws)
    if "sku" not in col_map:
        log.error("Could not find the SKU/article column — cannot match products. Aborting.")
        sys.exit(1)

    log.info("Downloading and merging supplier feeds...")
    offers = feeds.collect_offers()
    if not offers:
        log.error("No products could be collected from any supplier feed. Aborting — "
                   "the base file was left untouched.")
        sys.exit(1)
    log.info("Collected fresh data for %d products from suppliers.", len(offers))

    sku_col = col_map["sku"]
    # mode="stock" -> only touch availability/quantity, price stays exactly as in
    # the base export. mode="price" -> only touch price, availability/quantity
    # stay as in the base export. mode="both" -> touch all three.
    price_col = col_map.get("price") if mode in ("price", "both") else None
    avail_col = col_map.get("available") if mode in ("stock", "both") else None
    qty_col = col_map.get("quantity") if mode in ("stock", "both") else None

    matched_in_base = set()
    updated_rows = 0

    for row in range(2, ws.max_row + 1):
        sku_cell = ws.cell(row=row, column=sku_col)
        sku = sku_cell.value
        if not sku:
            continue
        sku = str(sku).strip()
        offer = offers.get(sku)
        if offer is None:
            continue  # this existing product has no fresh data from suppliers today — leave it untouched

        matched_in_base.add(sku)
        changed = False

        if price_col and offer.price:
            try:
                ws.cell(row=row, column=price_col).value = float(feeds.normalize_price(offer.price))
                changed = True
            except ValueError:
                pass

        if avail_col and offer.available is not None:
            ws.cell(row=row, column=avail_col).value = "+" if offer.available else "-"
            changed = True

        if qty_col and offer.quantity is not None:
            try:
                ws.cell(row=row, column=qty_col).value = int(feeds.normalize_qty(offer.quantity))
                changed = True
            except ValueError:
                pass

        if changed:
            updated_rows += 1

    not_in_base = set(offers.keys()) - matched_in_base
    field_desc = {"stock": "availability/quantity", "price": "price", "both": "price/availability/quantity"}[mode]
    log.info("Updated %d existing product(s) with fresh %s.", updated_rows, field_desc)
    if not_in_base:
        log.warning(
            "%d SKU(s) from the supplier feeds have no matching row in the base "
            "export (i.e. they don't exist on the site yet) and were skipped, as "
            "requested — only existing products are updated. Example SKUs: %s",
            len(not_in_base), sorted(not_in_base)[:10],
        )

    wb.save(output_path)
    log.info("Done. Import-ready file written to %s", output_path.resolve())


def inspect_base(base_path: Path) -> None:
    wb = openpyxl.load_workbook(base_path, data_only=True)
    print(f"Sheets found: {wb.sheetnames}")
    if SHEET_NAME not in wb.sheetnames:
        print(f"!! Sheet '{SHEET_NAME}' not found.")
        return
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    print(f"\nTotal columns: {len(headers)}, total data rows: {ws.max_row - 1}\n")
    col_map = load_column_map(ws)
    print("Detected columns:")
    for field, idx in col_map.items():
        print(f"  {field:10s} -> column {idx} ('{headers[idx-1]}')")
    missing = set(HEADER_VARIANTS) - set(col_map)
    if missing:
        print(f"\nNOT FOUND: {missing} — check the exact header names in your file "
              f"and add them to HEADER_VARIANTS in this script if needed.")
    print("\nFirst data row sample for the detected columns:")
    for field, idx in col_map.items():
        print(f"  {field}: {ws.cell(row=2, column=idx).value!r}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", required=True, help="Path to the full Prom.ua export (.xlsx).")
    parser.add_argument("--output", default="prom_import.xlsx", help="Path to write the updated import file to.")
    parser.add_argument(
        "--mode", choices=["stock", "price", "both"], default="both",
        help="stock — update only Наявність/Кількість, leave price exactly as in "
             "the base export; price — update only Ціна, leave availability/"
             "quantity as in the base export; both — update all three "
             "(default).",
    )
    parser.add_argument("--inspect", action="store_true",
                         help="Just show which columns were detected in --base and exit.")
    args = parser.parse_args()

    base_path = Path(args.base)
    if not base_path.exists():
        log.error("Base file not found: %s", base_path)
        return 1

    if args.inspect:
        inspect_base(base_path)
        return 0

    update_workbook(base_path, Path(args.output), mode=args.mode)
    return 0


if __name__ == "__main__":
    sys.exit(main())
